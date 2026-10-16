import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference

# ============ 配置 ============
file1 = r'e:\Users\61402\xwechat_files\wxid_yn56hzya6mu622_70ce\msg\file\2026-10\21-26国家公务员最低进面分数线.xlsx'
file2 = r'e:\Users\61402\xwechat_files\wxid_yn56hzya6mu622_70ce\msg\file\2026-10\中央机关及其直属机构2027年度考试录用公务员招考简章.xls'
outfile = r'd:\py\list\国考数据分析报告.xlsx'

# 各年份分数列名映射
SCORE_COL = {
    '2026': '最低进面分数',
    '2025': '进面分',
    '2024': '分数线',
    '2023': '最低进面分',
    '2022': '入面分数',
    '2021': '最低进面分数',
}
YEARS = ['2021', '2022', '2023', '2024', '2025', '2026']

# 地区列名映射
REGION_COL = {
    '2026': '地区',
    '2025': '地区',
    '2024': '地区',
    '2023': '工作地点',
    '2022': '工作地点',
    '2021': '工作地点',
}

# ============ 读取历年数据 ============
print('读取历年数据...')
year_data = {}
for year in YEARS:
    if year == '2026':
        df = pd.read_excel(file1, sheet_name=year)
    else:
        df = pd.read_excel(file1, sheet_name=year, header=1)
    df['年份'] = int(year)
    # 重命名分数列为统一名称
    score_col = SCORE_COL[year]
    df['进面分数'] = pd.to_numeric(df[score_col], errors='coerce')
    # 地区列统一
    rcol = REGION_COL[year]
    df['地区'] = df[rcol].astype(str).str.strip()
    # 部门名称统一
    dcol = '招录机关' if '招录机关' in df.columns else '部门名称'
    df['部门名称'] = df[dcol].astype(str).str.strip()
    # 招考人数
    df['招考人数'] = pd.to_numeric(df.get('招考人数', 0), errors='coerce').fillna(0).astype(int)
    year_data[year] = df
    print(f'  {year}: {len(df)} 条')

# 合并所有年份
all_years = pd.concat(year_data.values(), ignore_index=True)
valid_scores = all_years[all_years['进面分数'].notna()].copy()
print(f'合并后共 {len(all_years)} 条，有效分数 {len(valid_scores)} 条')

# ============ 读取2027职位数据 ============
print('读取2027招考简章...')
xl2 = pd.ExcelFile(file2)
sheets_2027 = [s for s in xl2.sheet_names if s != 'Sheet1']
df_2027_list = []
for s in sheets_2027:
    df = pd.read_excel(file2, sheet_name=s, header=1)
    df['机构类别'] = s
    df_2027_list.append(df)
df_2027 = pd.concat(df_2027_list, ignore_index=True)
df_2027['部门名称'] = df_2027['部门名称'].astype(str).str.strip()
df_2027['招考人数'] = pd.to_numeric(df_2027.get('招考人数', 0), errors='coerce').fillna(0).astype(int)
print(f'  2027: {len(df_2027)} 条')

# 按部门计算历年平均进面分数（用于2027匹配）
dept_scores = valid_scores.groupby('部门名称')['进面分数'].agg(['mean', 'median', 'min', 'max', 'count']).reset_index()
dept_scores.columns = ['部门名称', '历年平均分', '历年中位数', '历年最低分', '历年最高分', '历年职位数']

# ============ 创建Excel工作簿 ============
wb = Workbook()

# 样式定义
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
HEADER_FILL = PatternFill('solid', fgColor='1890FF')
SUBHEADER_FONT = Font(name='Arial', bold=True, size=11)
KPI_FILL = PatternFill('solid', fgColor='EAF2FF')
ZEBRA1 = PatternFill('solid', fgColor='FFFFFF')
ZEBRA2 = PatternFill('solid', fgColor='F7F9FC')
THIN = Side(style='thin', color='D9DEE7')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)

def style_header(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

def style_data(ws, start_row, end_row, max_col):
    for r in range(start_row, end_row + 1):
        fill = ZEBRA1 if (r - start_row) % 2 == 0 else ZEBRA2
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            cell.fill = fill
            cell.border = BORDER
            cell.alignment = CENTER if c <= 3 else LEFT

def auto_width(ws, max_col, min_w=10, max_w=50):
    for c in range(1, max_col + 1):
        letter = get_column_letter(c)
        max_len = min_w
        for row in ws.iter_rows(min_col=c, max_col=c, values_only=True):
            v = row[0]
            if v is not None:
                # 中文按2字符宽度
                l = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
                max_len = min(max(max_len, l), max_w)
        ws.column_dimensions[letter].width = max_len + 2

# ===================== Sheet 1: 数据概览 =====================
ws1 = wb.active
ws1.title = '数据概览'
ws1['A1'] = '国家公务员考试数据分析报告'
ws1['A1'].font = Font(name='Arial', bold=True, size=16, color='1890FF')
ws1.merge_cells('A1:D1')

ws1['A3'] = '一、数据集概览'
ws1['A3'].font = SUBHEADER_FONT

overview = [
    ['数据集', '年份/类别', '职位数', '有效分数记录数'],
    ['历年进面分数线', '2021', len(year_data['2021']), int(year_data['2021']['进面分数'].notna().sum())],
    ['历年进面分数线', '2022', len(year_data['2022']), int(year_data['2022']['进面分数'].notna().sum())],
    ['历年进面分数线', '2023', len(year_data['2023']), int(year_data['2023']['进面分数'].notna().sum())],
    ['历年进面分数线', '2024', len(year_data['2024']), int(year_data['2024']['进面分数'].notna().sum())],
    ['历年进面分数线', '2025', len(year_data['2025']), int(year_data['2025']['进面分数'].notna().sum())],
    ['历年进面分数线', '2026', len(year_data['2026']), int(year_data['2026']['进面分数'].notna().sum())],
    ['2027招考简章', '全部机构', len(df_2027), '-'],
]
for i, row in enumerate(overview):
    for j, val in enumerate(row):
        ws1.cell(4 + i, 1 + j, val)
style_header(ws1, 4, 4)
style_data(ws1, 5, 4 + len(overview) - 1, 4)

# 关键KPI
kpi_start = 4 + len(overview) + 2
ws1.cell(kpi_start, 1, '二、关键指标').font = SUBHEADER_FONT

total_2027 = len(df_2027)
total_recruit_2027 = int(df_2027['招考人数'].sum())
avg_score_all = valid_scores['进面分数'].mean()
max_score = valid_scores['进面分数'].max()
min_score = valid_scores['进面分数'].min()

kpis = [
    ['指标', '数值', '说明'],
    ['2027总职位数', total_2027, '2027年度国考招考职位总数'],
    ['2027总招考人数', total_recruit_2027, '2027年度国考计划招录总人数'],
    ['历年平均进面分', f'{avg_score_all:.1f}', '2021-2026所有职位平均最低进面分数'],
    ['历年最高进面分', f'{max_score:.1f}', '2021-2026职位最高最低进面分数'],
    ['历年最低进面分', f'{min_score:.1f}', '2021-2026职位最低最低进面分数'],
]
for i, row in enumerate(kpis):
    for j, val in enumerate(row):
        ws1.cell(kpi_start + 1 + i, 1 + j, val)
style_header(ws1, kpi_start + 1, 3)
style_data(ws1, kpi_start + 2, kpi_start + 1 + len(kpis) - 1, 3)
# 高亮KPI数值列
for r in range(kpi_start + 2, kpi_start + 2 + len(kpis) - 1):
    ws1.cell(r, 2).fill = KPI_FILL
    ws1.cell(r, 2).font = Font(name='Arial', bold=True, color='1890FF', size=12)

auto_width(ws1, 4)

# ===================== Sheet 2: 历年分数趋势 =====================
ws2 = wb.create_sheet('历年分数趋势')
trend_cols = ['年份', '职位数', '有效分数数', '最低分', '最高分', '平均分', '中位数', '平均分同比']
for j, c in enumerate(trend_cols):
    ws2.cell(1, j + 1, c)

prev_avg = None
for i, year in enumerate(YEARS):
    df = year_data[year]
    scores = df['进面分数'].dropna()
    avg = scores.mean()
    row = [
        int(year),
        len(df),
        len(scores),
        round(scores.min(), 1),
        round(scores.max(), 1),
        round(avg, 1),
        round(scores.median(), 1),
    ]
    if prev_avg is not None:
        row.append(round((avg - prev_avg) / prev_avg * 100, 1))
    else:
        row.append(None)
    for j, v in enumerate(row):
        ws2.cell(i + 2, j + 1, v)
    prev_avg = avg

style_header(ws2, 1, len(trend_cols))
style_data(ws2, 2, len(YEARS) + 1, len(trend_cols))

# 添加趋势图
chart = LineChart()
chart.title = '历年平均进面分数趋势'
chart.y_axis.title = '平均分'
chart.x_axis.title = '年份'
data = Reference(ws2, min_col=6, min_row=1, max_row=len(YEARS) + 1)
cats = Reference(ws2, min_col=1, min_row=2, max_row=len(YEARS) + 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 12
chart.width = 18
ws2.add_chart(chart, f'A{len(YEARS) + 4}')

auto_width(ws2, len(trend_cols))

# ===================== Sheet 3: 地区分数对比 =====================
ws3 = wb.create_sheet('地区分数对比')
# 取主要地区（按2026职位数排序前15）
top_regions = year_data['2026']['地区'].value_counts().head(15).index.tolist()

region_cols = ['地区'] + [f'{y}平均分' for y in YEARS] + ['六年平均']
for j, c in enumerate(region_cols):
    ws3.cell(1, j + 1, c)

for i, region in enumerate(top_regions):
    row = [region]
    row_scores = []
    for year in YEARS:
        df = year_data[year]
        s = df[df['地区'] == region]['进面分数'].dropna()
        if len(s) > 0:
            row.append(round(s.mean(), 1))
            row_scores.append(s.mean())
        else:
            row.append(None)
    if row_scores:
        row.append(round(np.mean(row_scores), 1))
    else:
        row.append(None)
    for j, v in enumerate(row):
        ws3.cell(i + 2, j + 1, v)

style_header(ws3, 1, len(region_cols))
style_data(ws3, 2, len(top_regions) + 1, len(region_cols))
auto_width(ws3, len(region_cols))

# ===================== Sheet 4: 2027职位分布统计 =====================
ws4 = wb.create_sheet('2027职位分布统计')

# 4.1 按地区
ws4['A1'] = '一、按工作地区分布'
ws4['A1'].font = SUBHEADER_FONT
region_stats = df_2027.groupby('工作地点').agg(职位数=('招考人数', 'count'), 招考人数=('招考人数', 'sum')).reset_index()
region_stats.columns = ['工作地区', '职位数', '招考人数']
region_stats = region_stats.sort_values('招考人数', ascending=False).head(20)
ws4.append(['工作地区', '职位数', '招考人数'])
for _, r in region_stats.iterrows():
    ws4.append([r['工作地区'], int(r['职位数']), int(r['招考人数'])])
r1_start = 2
r1_end = r1_start + len(region_stats) - 1
style_header(ws4, r1_start, 3)
style_data(ws4, r1_start + 1, r1_end, 3)

# 4.2 按学历
r2_start = r1_end + 3
ws4.cell(r2_start - 1, 1, '二、按学历要求分布').font = SUBHEADER_FONT
edu_stats = df_2027.groupby('学历').agg(职位数=('招考人数', 'count'), 招考人数=('招考人数', 'sum')).reset_index()
edu_stats.columns = ['学历', '职位数', '招考人数']
edu_stats = edu_stats.sort_values('招考人数', ascending=False)
ws4.cell(r2_start, 1, '学历')
ws4.cell(r2_start, 2, '职位数')
ws4.cell(r2_start, 3, '招考人数')
for i, r in enumerate(edu_stats.itertuples(index=False)):
    ws4.cell(r2_start + 1 + i, 1, r.学历)
    ws4.cell(r2_start + 1 + i, 2, int(r.职位数))
    ws4.cell(r2_start + 1 + i, 3, int(r.招考人数))
r2_end = r2_start + len(edu_stats)
style_header(ws4, r2_start, 3)
style_data(ws4, r2_start + 1, r2_end, 3)

# 4.3 按政治面貌
r3_start = r2_end + 3
ws4.cell(r3_start - 1, 1, '三、按政治面貌要求分布').font = SUBHEADER_FONT
pol_stats = df_2027.groupby('政治面貌').agg(职位数=('招考人数', 'count'), 招考人数=('招考人数', 'sum')).reset_index()
pol_stats.columns = ['政治面貌', '职位数', '招考人数']
pol_stats = pol_stats.sort_values('招考人数', ascending=False)
ws4.cell(r3_start, 1, '政治面貌')
ws4.cell(r3_start, 2, '职位数')
ws4.cell(r3_start, 3, '招考人数')
for i, r in enumerate(pol_stats.itertuples(index=False)):
    ws4.cell(r3_start + 1 + i, 1, r.政治面貌)
    ws4.cell(r3_start + 1 + i, 2, int(r.职位数))
    ws4.cell(r3_start + 1 + i, 3, int(r.招考人数))
r3_end = r3_start + len(pol_stats)
style_header(ws4, r3_start, 3)
style_data(ws4, r3_start + 1, r3_end, 3)

# 4.4 按机构类别
r4_start = r3_end + 3
ws4.cell(r4_start - 1, 1, '四、按机构类别分布').font = SUBHEADER_FONT
inst_stats = df_2027.groupby('机构类别').agg(职位数=('招考人数', 'count'), 招考人数=('招考人数', 'sum')).reset_index()
inst_stats.columns = ['机构类别', '职位数', '招考人数']
inst_stats = inst_stats.sort_values('招考人数', ascending=False)
ws4.cell(r4_start, 1, '机构类别')
ws4.cell(r4_start, 2, '职位数')
ws4.cell(r4_start, 3, '招考人数')
for i, r in enumerate(inst_stats.itertuples(index=False)):
    ws4.cell(r4_start + 1 + i, 1, r.机构类别)
    ws4.cell(r4_start + 1 + i, 2, int(r.职位数))
    ws4.cell(r4_start + 1 + i, 3, int(r.招考人数))
r4_end = r4_start + len(inst_stats)
style_header(ws4, r4_start, 3)
style_data(ws4, r4_start + 1, r4_end, 3)

auto_width(ws4, 3, min_w=12)

# ===================== Sheet 5: 分数匹配参考 =====================
ws5 = wb.create_sheet('分数匹配参考')
# 2027职位按部门匹配历年平均分
match_cols = ['部门名称', '2027职位数', '2027招考人数', '历年职位数', '历年平均分', '历年中位数', '历年最低分', '历年最高分']
for j, c in enumerate(match_cols):
    ws5.cell(1, j + 1, c)

# 2027按部门统计
dept_2027 = df_2027.groupby('部门名称').agg(职位数=('招考人数', 'count'), 招考人数=('招考人数', 'sum')).reset_index()
dept_2027.columns = ['部门名称', '2027职位数', '2027招考人数']
# 合并历年分数
merged = dept_2027.merge(dept_scores, on='部门名称', how='left')
merged = merged.sort_values('2027招考人数', ascending=False).head(100)

for i, (_, r) in enumerate(merged.iterrows()):
    row = [
        r['部门名称'],
        int(r['2027职位数']) if pd.notna(r['2027职位数']) else 0,
        int(r['2027招考人数']) if pd.notna(r['2027招考人数']) else 0,
        int(r['历年职位数']) if pd.notna(r['历年职位数']) else 0,
        round(r['历年平均分'], 1) if pd.notna(r['历年平均分']) else None,
        round(r['历年中位数'], 1) if pd.notna(r['历年中位数']) else None,
        round(r['历年最低分'], 1) if pd.notna(r['历年最低分']) else None,
        round(r['历年最高分'], 1) if pd.notna(r['历年最高分']) else None,
    ]
    for j, v in enumerate(row):
        ws5.cell(i + 2, j + 1, v)

style_header(ws5, 1, len(match_cols))
style_data(ws5, 2, len(merged) + 1, len(match_cols))
auto_width(ws5, len(match_cols))

# ===================== Sheet 6: 2027职位明细 =====================
ws6 = wb.create_sheet('2027职位明细')
detail_cols = ['部门名称', '用人司局', '招考职位', '职位代码', '机构性质', '招考人数', '工作地点', '学历', '政治面貌', '专业', '历年平均分']
for j, c in enumerate(detail_cols):
    ws6.cell(1, j + 1, c)

# 合并分数
detail = df_2027.merge(dept_scores[['部门名称', '历年平均分']], on='部门名称', how='left')
detail = detail.sort_values(['部门名称', '招考人数'], ascending=[True, False])
# 限制输出条数以控制文件大小
MAX_ROWS = 5000
detail_sample = detail.head(MAX_ROWS)
for i, r in enumerate(detail_sample.itertuples(index=False)):
    row = [
        getattr(r, '部门名称', ''),
        getattr(r, '用人司局', ''),
        getattr(r, '招考职位', ''),
        str(getattr(r, '职位代码', '')),
        getattr(r, '机构性质', ''),
        int(getattr(r, '招考人数', 0)) if pd.notna(getattr(r, '招考人数', 0)) else 0,
        getattr(r, '工作地点', ''),
        getattr(r, '学历', ''),
        getattr(r, '政治面貌', ''),
        str(getattr(r, '专业', ''))[:200],
        round(getattr(r, '历年平均分', None), 1) if pd.notna(getattr(r, '历年平均分', None)) else None,
    ]
    for j, v in enumerate(row):
        ws6.cell(i + 2, j + 1, v)

style_header(ws6, 1, len(detail_cols))
style_data(ws6, 2, len(detail_sample) + 1, len(detail_cols))
auto_width(ws6, len(detail_cols), max_w=60)
ws6.freeze_panes = 'A2'

# 保存
wb.save(outfile)
print(f'\n报告已生成: {outfile}')
print(f'Sheet1 数据概览')
print(f'Sheet2 历年分数趋势 ({len(YEARS)}行)')
print(f'Sheet3 地区分数对比 ({len(top_regions)}地区)')
print(f'Sheet4 2027职位分布统计')
print(f'Sheet5 分数匹配参考 (前100部门)')
print(f'Sheet6 2027职位明细 (前{MAX_ROWS}条)')
